import openpyxl
from openpyxl.styles import PatternFill
from openpyxl_simple.styles import sample_style_func, write_cell
from openpyxl_simple.writer import write_xlsx_ll


def test_style_func_applied():
    wb = openpyxl.Workbook()
    ws = wb.active
    custom_fill = PatternFill(patternType="solid", fgColor="FF0000")

    def custom_styler(cell):
        if cell.column == 2:
            cell.fill = custom_fill

    write_cell(ws, 1, 1, "A1", style_func=custom_styler)
    write_cell(ws, 1, 2, "B1", style_func=custom_styler)

    assert ws.cell(row=1, column=1).fill.fill_type is None
    assert ws.cell(row=1, column=2).fill.fill_type == "solid"
    assert ws.cell(row=1, column=2).fill.fgColor.rgb == "00FF0000"


def test_sample_style_func():
    wb = openpyxl.Workbook()
    ws = wb.active

    write_cell(ws, 1, 1, "row 1", style_func=sample_style_func)
    write_cell(ws, 2, 1, "row 2", style_func=sample_style_func)

    assert ws.cell(row=1, column=1).fill.fill_type is None
    assert ws.cell(row=2, column=1).fill.fill_type == "solid"
    assert ws.cell(row=2, column=1).fill.fgColor.rgb == "00999999"


def test_write_xlsx_with_style_func(tmp_path):
    fpath = tmp_path / "styled.xlsx"
    data = [["h1", "h2"], ["v1", "v2"], ["v3", "v4"]]
    write_xlsx_ll(str(fpath), data, style_func=sample_style_func)

    wb = openpyxl.load_workbook(str(fpath))
    ws = wb.active
    assert ws.cell(row=1, column=1).fill.fill_type is None
    assert ws.cell(row=2, column=1).fill.fill_type == "solid"
    assert ws.cell(row=3, column=1).fill.fill_type is None
    wb.close()
