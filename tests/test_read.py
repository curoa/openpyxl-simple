import csv
import openpyxl
from openpyxl_simple.reader import (
    load_as_dict,
    load_as_list,
    load_csv_as_dict,
    load_csv_as_list,
    load_ws_as_dict,
    load_ws_as_list,
    load_xlsx_as_dict,
    load_xlsx_as_list,
)


def test_read_ws_as_list():
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["A", "B", "C"])
    ws.append([1, 2, 3])
    ws.append(["x", "y", "z"])

    result = load_ws_as_list(ws)
    assert result == [["A", "B", "C"], [1, 2, 3], ["x", "y", "z"]]


def test_read_ws_as_dict():
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["id", "name", "age"])
    ws.append([1, "Alice", 20])
    ws.append([2, "Bob", 30])

    result = load_ws_as_dict(ws)
    assert result == [
        {"id": 1, "name": "Alice", "age": 20},
        {"id": 2, "name": "Bob", "age": 30},
    ]


def test_read_xlsx(tmp_path):
    fpath = tmp_path / "test.xlsx"
    wb = openpyxl.Workbook()
    ws1 = wb.active
    ws1.title = "Sheet1"
    ws1.append(["col1", "col2"])
    ws1.append(["val1", "val2"])

    ws2 = wb.create_sheet(title="CustomSheet")
    ws2.append(["name", "score"])
    ws2.append(["charlie", 95])
    wb.save(fpath)
    wb.close()

    assert load_xlsx_as_list(str(fpath)) == [["col1", "col2"], ["val1", "val2"]]
    assert load_xlsx_as_dict(str(fpath)) == [{"col1": "val1", "col2": "val2"}]
    assert load_xlsx_as_list(str(fpath), sheet_name="CustomSheet") == [["name", "score"], ["charlie", 95]]
    assert load_xlsx_as_dict(str(fpath), sheet_name="CustomSheet") == [{"name": "charlie", "score": 95}]


def test_read_csv(tmp_path):
    fpath = tmp_path / "test.csv"
    with open(fpath, "w", encoding="utf-16", newline="") as f:
        writer = csv.writer(f)
        writer.writerow(["id", "title"])
        writer.writerow(["10", "itemA"])
        writer.writerow(["20", "itemB"])

    assert load_csv_as_list(str(fpath), encoding="utf-16") == [["id", "title"], ["10", "itemA"], ["20", "itemB"]]
    assert load_csv_as_dict(str(fpath), encoding="utf-16") == [
        {"id": "10", "title": "itemA"},
        {"id": "20", "title": "itemB"},
    ]


def test_read_dispatcher_fallback(tmp_path):
    csv_file = tmp_path / "fallback.csv"
    with open(csv_file, "w", encoding="utf-16", newline="") as f:
        writer = csv.writer(f)
        writer.writerow(["c1", "c2"])
        writer.writerow(["v1", "v2"])

    non_exist_xlsx = tmp_path / "fallback.xlsx"
    assert load_as_list(str(non_exist_xlsx)) == [["c1", "c2"], ["v1", "v2"]]

    csv_dict_file = tmp_path / "fallback_dict.csv"
    with open(csv_dict_file, "w", encoding="utf-8", newline="") as f:
        writer = csv.DictWriter(f, fieldnames=["key", "value"])
        writer.writeheader()
        writer.writerow({"key": "val1", "value": "val2"})

    non_exist_xlsx_dict = tmp_path / "fallback_dict.xlsx"
    assert load_as_dict(str(non_exist_xlsx_dict)) == [{"key": "val1", "value": "val2"}]
