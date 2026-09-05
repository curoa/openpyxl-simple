import openpyxl
from openpyxl_simple.styles import write_cell
from openpyxl_simple.writer import write_xlsx_dict


def test_url_hyperlink_https_and_http():
    wb = openpyxl.Workbook()
    ws = wb.active

    write_cell(ws, 1, 1, "https://example.com/test", style_func=None)
    write_cell(ws, 1, 2, "http://example.org", style_func=None)
    write_cell(ws, 1, 3, "https://example.com/with space", style_func=None)
    write_cell(ws, 1, 4, "plain text", style_func=None)
    write_cell(ws, 1, 5, 12345, style_func=None)

    assert ws.cell(row=1, column=1).hyperlink.target == "https://example.com/test"
    assert ws.cell(row=1, column=2).hyperlink.target == "http://example.org"
    assert ws.cell(row=1, column=3).hyperlink is None
    assert ws.cell(row=1, column=4).hyperlink is None
    assert ws.cell(row=1, column=5).hyperlink is None


def test_url_roundtrip_xlsx(tmp_path):
    fpath = tmp_path / "url_test.xlsx"
    header = ["url", "title"]
    data = [
        {"url": "https://example.com", "title": "Example"},
        {"url": "http://openpyxl.readthedocs.io", "title": "Doc"},
    ]
    write_xlsx_dict(str(fpath), header, data)

    wb = openpyxl.load_workbook(str(fpath))
    ws = wb.active
    assert ws.cell(row=2, column=1).hyperlink.target == "https://example.com"
    assert ws.cell(row=3, column=1).hyperlink.target == "http://openpyxl.readthedocs.io"
    assert ws.cell(row=2, column=2).hyperlink is None
    wb.close()
