import datetime
import openpyxl
from openpyxl_simple.reader import (
    load_as_dict,
    load_as_list,
    load_csv_as_dict,
    load_csv_as_list,
    load_xlsx_as_dict,
    load_xlsx_as_list,
)
from openpyxl_simple.writer import (
    write_csv_dict,
    write_csv_ll,
    write_dict,
    write_ll,
    write_ws_dict,
    write_ws_ll,
    write_xlsx_dict,
    write_xlsx_ll,
)


def test_write_ws_ll():
    wb = openpyxl.Workbook()
    ws = wb.active
    data = [["A", "B"], [1, 2]]
    write_ws_ll(ws, data)
    assert ws.cell(row=1, column=1).value == "A"
    assert ws.cell(row=1, column=2).value == "B"
    assert ws.cell(row=2, column=1).value == 1
    assert ws.cell(row=2, column=2).value == 2


def test_write_ws_dict():
    wb = openpyxl.Workbook()
    ws = wb.active
    header = ["id", "name"]
    data = [{"id": 1, "name": "Alice"}, {"id": 2}]
    write_ws_dict(ws, header, data)
    assert ws.cell(row=1, column=1).value == "id"
    assert ws.cell(row=1, column=2).value == "name"
    assert ws.cell(row=2, column=1).value == 1
    assert ws.cell(row=2, column=2).value == "Alice"
    assert ws.cell(row=3, column=1).value == 2
    assert ws.cell(row=3, column=2).value is None


def test_write_xlsx_ll_and_dict(tmp_path):
    # write_xlsx_ll
    ll_path = tmp_path / "ll.xlsx"
    write_xlsx_ll(str(ll_path), [["header1", "header2"], ["v1", "v2"]])
    assert load_xlsx_as_list(str(ll_path)) == [["header1", "header2"], ["v1", "v2"]]

    # write_xlsx_dict default sheet
    dict_path = tmp_path / "dict.xlsx"
    header = ["k1", "k2"]
    data = [{"k1": "a", "k2": "b"}]
    write_xlsx_dict(str(dict_path), header, data)
    assert load_xlsx_as_dict(str(dict_path)) == [{"k1": "a", "k2": "b"}]

    # write_xlsx_dict custom sheet in existing file
    data2 = [{"k1": "c", "k2": "d"}]
    write_xlsx_dict(str(dict_path), header, data2, sheet_name="NewSheet")
    assert load_xlsx_as_dict(str(dict_path), sheet_name="NewSheet") == [{"k1": "c", "k2": "d"}]


def test_write_csv_ll_and_dict(tmp_path):
    ll_path = tmp_path / "ll.csv"
    write_csv_ll(str(ll_path), [["c1", "c2"], [10, 20]])
    assert load_csv_as_list(str(ll_path)) == [["c1", "c2"], ["10", "20"]]

    dict_path = tmp_path / "dict.csv"
    header = ["c1", "c2"]
    data = [{"c1": "foo", "c2": "bar"}]
    write_csv_dict(str(dict_path), header, data)
    assert load_csv_as_dict(str(dict_path)) == [{"c1": "foo", "c2": "bar"}]


def test_write_dispatcher(tmp_path):
    xlsx_ll = tmp_path / "out_ll.xlsx"
    write_ll(str(xlsx_ll), [["col"], [1]])
    assert load_as_list(str(xlsx_ll)) == [["col"], [1]]

    csv_ll = tmp_path / "out_ll.csv"
    write_ll(str(csv_ll), [["col"], [1]])
    assert load_as_list(str(csv_ll)) == [["col"], ["1"]]

    xlsx_dict = tmp_path / "out_dict.xlsx"
    write_dict(str(xlsx_dict), ["col"], [{"col": "val"}])
    assert load_as_dict(str(xlsx_dict)) == [{"col": "val"}]

    csv_dict = tmp_path / "out_dict.csv"
    write_dict(str(csv_dict), ["col"], [{"col": "val"}])
    assert load_as_dict(str(csv_dict)) == [{"col": "val"}]


def test_write_and_read_datetime(tmp_path):
    fpath = tmp_path / "datetime_test.xlsx"
    now = datetime.datetime(2026, 9, 5, 12, 30, 45)
    today = datetime.date(2026, 9, 5)

    data = [
        ["datetime", "date"],
        [now, today],
    ]
    write_xlsx_ll(str(fpath), data)

    result_list = load_xlsx_as_list(str(fpath))
    assert result_list[0] == ["datetime", "date"]
    assert result_list[1][0] == now
    assert result_list[1][1] == datetime.datetime(2026, 9, 5, 0, 0)

    dict_data = [{"datetime": now, "date": today}]
    dict_fpath = tmp_path / "datetime_dict.xlsx"
    write_xlsx_dict(str(dict_fpath), ["datetime", "date"], dict_data)

    result_dict = load_xlsx_as_dict(str(dict_fpath))
    assert result_dict[0]["datetime"] == now
    assert result_dict[0]["date"] == datetime.datetime(2026, 9, 5, 0, 0)
