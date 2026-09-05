import csv
from pathlib import Path
import openpyxl
from openpyxl.styles import PatternFill
import openpyxl_simple as ops
from openpyxl_simple.decorator import deco_fname_check, deco_to_path
from openpyxl_simple.encoding import guess_utf_encoding
from openpyxl_simple.reader import (
    load_as_dict,
    load_as_list,
    load_csv_as_dict,
    load_csv_as_list,
    load_ws_as_dict,
    load_ws_as_list,
    load_xlsx_as_dict,
    load_xlsx_as_list,
)
from openpyxl_simple.styles import sample_style_func, write_cell
from openpyxl_simple.table import collect_header, tie_key_value
from openpyxl_simple.writer import (
    write_csv_dict,
    write_csv_ll,
    write_dict,
    write_ll,
    write_ws_dict,
    write_ws_ll,
    write_xlsx_dict,
    write_xlsx_ll,
)

# read tests
def test_read_ws_as_list():
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["A", "B", "C"])
    ws.append([1, 2, 3])
    ws.append(["x", "y", "z"])

    result = load_ws_as_list(ws)
    assert result == [["A", "B", "C"], [1, 2, 3], ["x", "y", "z"]]

def test_read_ws_as_dict():
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["id", "name", "age"])
    ws.append([1, "Alice", 20])
    ws.append([2, "Bob", 30])

    result = load_ws_as_dict(ws)
    assert result == [
        {"id": 1, "name": "Alice", "age": 20},
        {"id": 2, "name": "Bob", "age": 30},
    ]

def test_read_xlsx(tmp_path):
    fpath = tmp_path / "test.xlsx"
    wb = openpyxl.Workbook()
    ws1 = wb.active
    ws1.title = "Sheet1"
    ws1.append(["col1", "col2"])
    ws1.append(["val1", "val2"])

    ws2 = wb.create_sheet(title="CustomSheet")
    ws2.append(["name", "score"])
    ws2.append(["charlie", 95])
    wb.save(fpath)
    wb.close()

    assert load_xlsx_as_list(str(fpath)) == [["col1", "col2"], ["val1", "val2"]]
    assert load_xlsx_as_dict(str(fpath)) == [{"col1": "val1", "col2": "val2"}]
    assert load_xlsx_as_list(str(fpath), sheet_name="CustomSheet") == [["name", "score"], ["charlie", 95]]
    assert load_xlsx_as_dict(str(fpath), sheet_name="CustomSheet") == [{"name": "charlie", "score": 95}]

def test_read_csv(tmp_path):
    fpath = tmp_path / "test.csv"
    with open(fpath, "w", encoding="utf-16", newline="") as f:
        writer = csv.writer(f)
        writer.writerow(["id", "title"])
        writer.writerow(["10", "itemA"])
        writer.writerow(["20", "itemB"])

    assert load_csv_as_list(str(fpath), encoding="utf-16") == [["id", "title"], ["10", "itemA"], ["20", "itemB"]]
    assert load_csv_as_dict(str(fpath), encoding="utf-16") == [
        {"id": "10", "title": "itemA"},
        {"id": "20", "title": "itemB"},
    ]

def test_read_dispatcher_fallback(tmp_path):
    csv_file = tmp_path / "fallback.csv"
    with open(csv_file, "w", encoding="utf-16", newline="") as f:
        writer = csv.writer(f)
        writer.writerow(["c1", "c2"])
        writer.writerow(["v1", "v2"])

    non_exist_xlsx = tmp_path / "fallback.xlsx"
    assert load_as_list(str(non_exist_xlsx)) == [["c1", "c2"], ["v1", "v2"]]

    csv_dict_file = tmp_path / "fallback_dict.csv"
    with open(csv_dict_file, "w", encoding="utf-8", newline="") as f:
        writer = csv.DictWriter(f, fieldnames=["key", "value"])
        writer.writeheader()
        writer.writerow({"key": "val1", "value": "val2"})

    non_exist_xlsx_dict = tmp_path / "fallback_dict.xlsx"
    assert load_as_dict(str(non_exist_xlsx_dict)) == [{"key": "val1", "value": "val2"}]

# write tests
def test_write_ws_ll():
    wb = openpyxl.Workbook()
    ws = wb.active
    data = [["A", "B"], [1, 2]]
    write_ws_ll(ws, data)
    assert ws.cell(row=1, column=1).value == "A"
    assert ws.cell(row=1, column=2).value == "B"
    assert ws.cell(row=2, column=1).value == 1
    assert ws.cell(row=2, column=2).value == 2

def test_write_ws_dict():
    wb = openpyxl.Workbook()
    ws = wb.active
    header = ["id", "name"]
    data = [{"id": 1, "name": "Alice"}, {"id": 2}]
    write_ws_dict(ws, header, data)
    assert ws.cell(row=1, column=1).value == "id"
    assert ws.cell(row=1, column=2).value == "name"
    assert ws.cell(row=2, column=1).value == 1
    assert ws.cell(row=2, column=2).value == "Alice"
    assert ws.cell(row=3, column=1).value == 2
    assert ws.cell(row=3, column=2).value is None

def test_write_xlsx_ll_and_dict(tmp_path):
    # write_xlsx_ll
    ll_path = tmp_path / "ll.xlsx"
    write_xlsx_ll(str(ll_path), [["header1", "header2"], ["v1", "v2"]])
    assert load_xlsx_as_list(str(ll_path)) == [["header1", "header2"], ["v1", "v2"]]

    # write_xlsx_dict default sheet
    dict_path = tmp_path / "dict.xlsx"
    header = ["k1", "k2"]
    data = [{"k1": "a", "k2": "b"}]
    write_xlsx_dict(str(dict_path), header, data)
    assert load_xlsx_as_dict(str(dict_path)) == [{"k1": "a", "k2": "b"}]

    # write_xlsx_dict custom sheet in existing file
    data2 = [{"k1": "c", "k2": "d"}]
    write_xlsx_dict(str(dict_path), header, data2, sheet_name="NewSheet")
    assert load_xlsx_as_dict(str(dict_path), sheet_name="NewSheet") == [{"k1": "c", "k2": "d"}]

def test_write_csv_ll_and_dict(tmp_path):
    ll_path = tmp_path / "ll.csv"
    write_csv_ll(str(ll_path), [["c1", "c2"], [10, 20]])
    assert load_csv_as_list(str(ll_path)) == [["c1", "c2"], ["10", "20"]]

    dict_path = tmp_path / "dict.csv"
    header = ["c1", "c2"]
    data = [{"c1": "foo", "c2": "bar"}]
    write_csv_dict(str(dict_path), header, data)
    assert load_csv_as_dict(str(dict_path)) == [{"c1": "foo", "c2": "bar"}]

def test_write_dispatcher(tmp_path):
    xlsx_ll = tmp_path / "out_ll.xlsx"
    write_ll(str(xlsx_ll), [["col"], [1]])
    assert load_as_list(str(xlsx_ll)) == [["col"], [1]]

    csv_ll = tmp_path / "out_ll.csv"
    write_ll(str(csv_ll), [["col"], [1]])
    assert load_as_list(str(csv_ll)) == [["col"], ["1"]]

    xlsx_dict = tmp_path / "out_dict.xlsx"
    write_dict(str(xlsx_dict), ["col"], [{"col": "val"}])
    assert load_as_dict(str(xlsx_dict)) == [{"col": "val"}]

    csv_dict = tmp_path / "out_dict.csv"
    write_dict(str(csv_dict), ["col"], [{"col": "val"}])
    assert load_as_dict(str(csv_dict)) == [{"col": "val"}]

# url tests
def test_url_hyperlink_https_and_http():
    wb = openpyxl.Workbook()
    ws = wb.active

    write_cell(ws, 1, 1, "https://example.com/test", style_func=None)
    write_cell(ws, 1, 2, "http://example.org", style_func=None)
    write_cell(ws, 1, 3, "https://example.com/with space", style_func=None)
    write_cell(ws, 1, 4, "plain text", style_func=None)
    write_cell(ws, 1, 5, 12345, style_func=None)

    assert ws.cell(row=1, column=1).hyperlink.target == "https://example.com/test"
    assert ws.cell(row=1, column=2).hyperlink.target == "http://example.org"
    assert ws.cell(row=1, column=3).hyperlink is None
    assert ws.cell(row=1, column=4).hyperlink is None
    assert ws.cell(row=1, column=5).hyperlink is None

def test_url_roundtrip_xlsx(tmp_path):
    fpath = tmp_path / "url_test.xlsx"
    header = ["url", "title"]
    data = [
        {"url": "https://example.com", "title": "Example"},
        {"url": "http://openpyxl.readthedocs.io", "title": "Doc"},
    ]
    write_xlsx_dict(str(fpath), header, data)

    wb = openpyxl.load_workbook(str(fpath))
    ws = wb.active
    assert ws.cell(row=2, column=1).hyperlink.target == "https://example.com"
    assert ws.cell(row=3, column=1).hyperlink.target == "http://openpyxl.readthedocs.io"
    assert ws.cell(row=2, column=2).hyperlink is None
    wb.close()

# style tests
def test_style_func_applied():
    wb = openpyxl.Workbook()
    ws = wb.active
    custom_fill = PatternFill(patternType="solid", fgColor="FF0000")

    def custom_styler(cell):
        if cell.column == 2:
            cell.fill = custom_fill

    write_cell(ws, 1, 1, "A1", style_func=custom_styler)
    write_cell(ws, 1, 2, "B1", style_func=custom_styler)

    assert ws.cell(row=1, column=1).fill.fill_type is None
    assert ws.cell(row=1, column=2).fill.fill_type == "solid"
    assert ws.cell(row=1, column=2).fill.fgColor.rgb == "00FF0000"

def test_sample_style_func():
    wb = openpyxl.Workbook()
    ws = wb.active

    write_cell(ws, 1, 1, "row 1", style_func=sample_style_func)
    write_cell(ws, 2, 1, "row 2", style_func=sample_style_func)

    assert ws.cell(row=1, column=1).fill.fill_type is None
    assert ws.cell(row=2, column=1).fill.fill_type == "solid"
    assert ws.cell(row=2, column=1).fill.fgColor.rgb == "00999999"

def test_write_xlsx_with_style_func(tmp_path):
    fpath = tmp_path / "styled.xlsx"
    data = [["h1", "h2"], ["v1", "v2"], ["v3", "v4"]]
    write_xlsx_ll(str(fpath), data, style_func=sample_style_func)

    wb = openpyxl.load_workbook(str(fpath))
    ws = wb.active
    assert ws.cell(row=1, column=1).fill.fill_type is None
    assert ws.cell(row=2, column=1).fill.fill_type == "solid"
    assert ws.cell(row=3, column=1).fill.fill_type is None
    wb.close()

# common / existing tests
def test_exports():
    for name in ops.__all__:
        assert hasattr(ops, name), f"Missing {name}"

def test_deco_to_path():
    @deco_to_path
    def dummy_fpath(fpath):
        return fpath

    assert isinstance(dummy_fpath("foo.txt"), type(Path("foo.txt")))
    assert isinstance(dummy_fpath(fpath="foo.txt"), type(Path("foo.txt")))
    assert isinstance(dummy_fpath(Path("foo.txt")), type(Path("foo.txt")))

    @deco_to_path(arg_name="filepath")
    def dummy_filepath(filepath):
        return filepath

    assert isinstance(dummy_filepath("bar.txt"), type(Path("bar.txt")))
    assert isinstance(dummy_filepath(filepath="bar.txt"), type(Path("bar.txt")))

    @deco_to_path
    def dummy_default(fpath="default.txt"):
        return fpath

    assert isinstance(dummy_default(), type(Path("default.txt")))
    assert dummy_default() == Path("default.txt")

def test_deco_fname_check():
    @deco_fname_check("xlsx")
    def dummy_save(fpath):
        return True

    assert dummy_save("test.xlsx") is True
    assert dummy_save("test.csv") is True

def test_encoding(tmp_path):
    f = tmp_path / "test_utf8.txt"
    f.write_text("hello", encoding="utf-8")
    assert guess_utf_encoding(str(f)) == "utf-8"

    f_bom = tmp_path / "test_bom.txt"
    f_bom.write_bytes(b"\xef\xbb\xbfhello")
    assert guess_utf_encoding(str(f_bom)) == "utf-8-sig"

def test_table():
    header = ["id", "name"]
    rows = [[1, "Alice"], [2, "Bob"]]
    data = tie_key_value(header, rows)
    assert data == [{"id": 1, "name": "Alice"}, {"id": 2, "name": "Bob"}]
    collected = collect_header(data)
    assert collected == ["id", "name"]

