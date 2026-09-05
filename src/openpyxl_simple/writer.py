import csv

import openpyxl

from openpyxl_simple.styles import write_cell
from openpyxl_simple.utils import deco_fname_check, to_path

def write_ll(fpath, data):
    fpath = to_path(fpath)
    if fpath.suffix == ".csv":
        write_csv_ll(fpath, data)
    else:
        write_xlsx_ll(fpath, data)

def write_dict(fpath, header, data):
    fpath = to_path(fpath)
    if fpath.suffix == ".csv":
        write_csv_dict(fpath, header, data)
    else:
        write_xlsx_dict(fpath, header, data)

@deco_fname_check("xlsx")
def write_xlsx_ll(fpath, data, style_func=None):
    wb = openpyxl.Workbook()
    ws = wb.active
    write_ws_ll(ws, data, style_func)
    wb.save(fpath)
    wb.close()

def write_ws_ll(ws, data, style_func=None):
    for i, row in enumerate(data, 1):
        for j, value in enumerate(row, 1):
            write_cell(ws, i, j, value, style_func)

@deco_fname_check("csv")
def write_csv_ll(fpath, data):
    writer = csv.writer(open(fpath, "w", encoding="utf-16"))
    writer.writerows(data)

# data: list of dict
@deco_fname_check("xlsx")
def write_xlsx_dict(fpath, header, data, sheet_name=None, style_func=None):
    fpath = to_path(fpath)
    if sheet_name is None:
        wb = openpyxl.Workbook()
        ws = wb.active
    else:
        if fpath.exists():
            wb = openpyxl.load_workbook(fpath)
        else:
            wb = openpyxl.Workbook()
        if sheet_name in wb.sheetnames:
            ws = wb[sheet_name]
        else:
            ws = wb.create_sheet(title=sheet_name)
    write_ws_dict(ws, header, data, style_func)
    wb.save(fpath)
    wb.close()

# data: list of dict
def write_ws_dict(ws, header, data, style_func=None):
    for j, key in enumerate(header, 1):
        ws.cell(row=1, column=j, value=key)
    for i, row in enumerate(data, 2):
        for j, key in enumerate(header, 1):
            value = row.get(key, None)
            write_cell(ws, i, j, value, style_func)

# use `kwargs={extrasaction: "ignore"}`
@deco_fname_check("csv")
def write_csv_dict(fpath, header, data, **kwargs):
    fpath = to_path(fpath)
    writer = csv.DictWriter(open(fpath, "w", encoding="utf-16"), header, **kwargs)
    writer.writeheader()
    for row in data:
        writer.writerow(row)
