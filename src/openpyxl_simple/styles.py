from openpyxl.styles import PatternFill

ptn_fill_gray = PatternFill(patternType="solid", fgColor="999999")

def sample_style_func(cell):
    # ref. https://openpyxl.readthedocs.io/en/stable/api/openpyxl.cell.cell.html
    if cell.row % 2 == 0:
        cell.fill = ptn_fill_gray

def write_cell(ws, i, j, value, style_func):
    cell = ws.cell(row=i, column=j, value=value)
    call_link = "https://call.ctrlq.org/"
    if type(value) is str and value.startswith(call_link):
        cell.hyperlink = value
        cell.value = value[len(call_link):]
    elif type(value) is str and (value.startswith("https://") or value.startswith("http://")) and " " not in value:
        cell.hyperlink = value
    if style_func is not None:
        style_func(cell)
